"""
Update all uploaded products to use real Size + Color variants from the xlsx.

Run after upload_products.py has finished.

Cleans:
  - "COMFORT COLORS® - S" → "S"
  - "Unisex - XL"         → "XL"
  - "SWEATSHIRT - M"      → "M"
  - "Kids Tee - L"        → "L"
  - "Stickers (6&quot;)"  → skipped
  - "Charcoal : Pepper"   → two colours: Charcoal, Pepper
  - "Other Colors"        → skipped

Variant grid: Size × Color (capped at 100 per Shopify's limit).
Products with no COLORS column get Size-only variants.
"""
import csv, json, os, re, sys, time
import openpyxl
import httpx
from dotenv import load_dotenv

sys.stdout.reconfigure(encoding="utf-8")
load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))

STORE = os.environ["SHOPIFY_STORE"]
H     = {"X-Shopify-Access-Token": os.environ["SHOPIFY_TOKEN"], "Content-Type": "application/json"}
API   = f"https://{STORE}/admin/api/2025-01"
BASE  = os.path.dirname(os.path.abspath(__file__))

XLSX_PATH     = os.path.join(BASE, "..", "scraper", "matcha_tees_catalogue.xlsx")
PRODUCTS_FILE = os.path.join(BASE, "products_uploaded.json")
PRICE         = "35.99"

SIZE_ORDER = ["XS", "S", "M", "L", "XL", "2XL", "3XL", "4XL", "5XL"]

SIZE_PREFIXES = [
    "COMFORT COLORS® - ",
    "COMFORT COLORS® -",
    "Unisex - ",
    "SWEATSHIRT - ",
    "Kids Tee - ",
]

SKIP_COLORS = {"Other Colors", ""}


# ── parsers ───────────────────────────────────────────────────────────────────

def parse_sizes(raw):
    seen, result = set(), []
    for token in raw.split(","):
        t = token.strip()
        for pfx in SIZE_PREFIXES:
            if t.startswith(pfx):
                t = t[len(pfx):].strip()
                break
        if not t or t.lower().startswith("sticker"):
            continue
        if t not in seen:
            seen.add(t)
            result.append(t)

    def size_key(s):
        try:
            return SIZE_ORDER.index(s)
        except ValueError:
            return len(SIZE_ORDER)

    return sorted(result, key=size_key)


def parse_colors(raw):
    if not raw.strip():
        return []
    seen, result = set(), []
    for part in raw.split(","):
        for sub in part.split(" : "):
            c = sub.strip()
            if c and c not in SKIP_COLORS and c not in seen:
                seen.add(c)
                result.append(c)
    return result


# ── shopify ───────────────────────────────────────────────────────────────────

def build_payload(product_id, sizes, colors):
    if colors:
        # cap to stay under 100
        max_colors = max(1, 100 // len(sizes))
        colors = colors[:max_colors]
        variants = [
            {"option1": sz, "option2": cl,
             "price": PRICE, "requires_shipping": True, "taxable": True}
            for sz in sizes for cl in colors
        ]
        options = [
            {"name": "Size",  "values": sizes},
            {"name": "Color", "values": colors},
        ]
    else:
        variants = [
            {"option1": sz, "price": PRICE,
             "requires_shipping": True, "taxable": True}
            for sz in sizes
        ]
        options = [{"name": "Size", "values": sizes}]

    return {"product": {"id": product_id, "options": options, "variants": variants}}


def update_product(product_id, sizes, colors):
    payload = build_payload(product_id, sizes, colors)
    for attempt in range(3):
        r = httpx.put(f"{API}/products/{product_id}.json",
                      headers=H, json=payload, timeout=60)
        if r.status_code == 429:
            wait = int(r.headers.get("Retry-After", 10))
            print(f"    Rate limited — waiting {wait}s")
            time.sleep(wait)
            continue
        return r.status_code, r.text[:300]
    return 0, "max retries"


# ── load catalogue ────────────────────────────────────────────────────────────

def load_catalogue():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c or "").strip()
               for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    catalogue = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        d = {headers[i]: str(v or "").strip() for i, v in enumerate(row) if i < len(headers)}
        if d.get("INCLUDE") != "True":
            continue
        raw_title = d.get("TITLE", "")
        clean = raw_title.split(",")[0].strip()[:100]
        catalogue[clean] = {
            "sizes":  parse_sizes(d.get("SIZES", "")),
            "colors": parse_colors(d.get("COLORS", "")),
        }
    wb.close()
    return catalogue


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    with open(PRODUCTS_FILE, encoding="utf-8") as f:
        uploaded = {k: v for k, v in json.load(f).items() if v}

    catalogue = load_catalogue()
    print(f"Uploaded products : {len(uploaded)}")
    print(f"Catalogue entries : {len(catalogue)}\n")

    ok = fail = skipped = 0

    for title, pid in uploaded.items():
        info = catalogue.get(title)
        if not info:
            print(f"  ? {title[:55]} — not found in catalogue, skipping")
            skipped += 1
            continue

        sizes  = info["sizes"]
        colors = info["colors"]

        if not sizes:
            sizes = ["S", "M", "L", "XL", "2XL", "3XL"]   # safe fallback

        variant_count = len(sizes) * max(len(colors), 1)
        color_label   = f"{len(colors)} colours" if colors else "no colour option"
        print(f"  {title[:50]}  →  {len(sizes)} sizes × {color_label} = {variant_count} variants")

        status, body = update_product(int(pid), sizes, colors)
        if status in (200, 201):
            ok += 1
            print(f"    ✓ updated (HTTP {status})")
        else:
            fail += 1
            print(f"    ✗ HTTP {status}: {body}")

        time.sleep(0.7)   # stay well under the 2-req/s REST limit

    print(f"\nFinished — {ok} updated  {fail} failed  {skipped} skipped")


if __name__ == "__main__":
    main()
